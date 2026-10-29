import hashlib
import json
import re

from docx import Document
from pypdf import PdfReader

from .config import DOCUMENTS


def corpus_fingerprint(directory=DOCUMENTS):
    digest = hashlib.sha256()
    for path in sorted(directory.iterdir()):
        if path.suffix.lower() in {'.pdf', '.docx', '.md', '.txt', '.json'}:
            digest.update(path.name.encode())
            digest.update(path.read_bytes())
    return digest.hexdigest()


def read_pages(path):
    if path.suffix.lower() == '.pdf':
        return [(i + 1, page.extract_text() or '') for i, page in enumerate(PdfReader(path).pages)]
    if path.suffix.lower() == '.docx':
        doc = Document(path)
        text = '\n'.join(p.text for p in doc.paragraphs)
        text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
        return [(None, text)]
    return [(None, path.read_text(encoding='utf-8'))]


def load_chunks(directory=DOCUMENTS):
    chunks = []
    for path in sorted(directory.iterdir()):
        if path.suffix.lower() not in {'.pdf', '.docx', '.md', '.txt'}:
            continue
        sidecar = path.with_suffix(path.suffix + '.json')
        metadata = json.loads(sidecar.read_text(encoding='utf-8')) if sidecar.exists() else {}
        if metadata.get('approved') is not True:
            raise ValueError(f'{path.name} : validation documentaire approved=true manquante.')
        section = 'Général'
        for page, text in read_pages(path):
            for paragraph in re.split(r'\n\s*\n', text):
                if paragraph.startswith('#'):
                    heading, _, paragraph = paragraph.partition('\n')
                    section = heading.lstrip('# ').strip()
                cleaned = re.sub(r'\s+', ' ', paragraph).strip()
                for start in range(0, len(cleaned), 850):
                    passage = cleaned[start:start + 1000]
                    if len(passage) < 30:
                        continue
                    chunks.append(dict(id=len(chunks) + 1, document=path.name,
                                       title=metadata.get('title', path.stem), section=section,
                                       page=page, url=metadata.get('url'), text=passage))
    return chunks
